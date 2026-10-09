import openpyxl

wb = openpyxl.load_workbook('Data Dashborad.xlsx', data_only=True)
print("Sheet names:", wb.sheetnames)

for name in wb.sheetnames:
    ws = wb[name]
    for r in range(1, min(10, ws.max_row + 1)):
        row_vals = [str(ws.cell(r, c).value) for c in range(1, min(25, ws.max_column + 1))]
        if any('CODE SBU' in v or 'Jan-26' in v or 'PENCAPAIAN' in v for v in row_vals):
            print(f"Sheet '{name}' row {r}:")
            for c in range(1, min(30, ws.max_column + 1)):
                val = ws.cell(r, c).value
                if val is not None:
                    print(f"  Col {c} ({openpyxl.utils.get_column_letter(c)}): {val}")
